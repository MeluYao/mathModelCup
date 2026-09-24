from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import load_workbook


@pytest.mark.skipif(
    not os.environ.get("D_Q1_TEMPLATE") or not os.environ.get("D_Q1_SUBMISSION"),
    reason="submission workbook paths are not configured",
)
def test_official_submission_preserves_template_and_covers_all_boxes() -> None:
    template_path = Path(os.environ["D_Q1_TEMPLATE"])
    output_path = Path(os.environ["D_Q1_SUBMISSION"])
    assert output_path.exists()

    original = load_workbook(template_path, data_only=False)
    result = load_workbook(output_path, data_only=False)
    assert result.sheetnames == original.sheetnames
    assert result.sheetnames == [
        "Q1_单点组批",
        "Q2_运输架次",
        "Q2_逐箱交付",
        "Q3_中继架次",
        "Q3_通信保障",
        "Q4_分区配置",
    ]

    for sheet_name in result.sheetnames[1:]:
        original_sheet = original[sheet_name]
        result_sheet = result[sheet_name]
        assert original_sheet.max_row == result_sheet.max_row
        assert original_sheet.max_column == result_sheet.max_column
        for row in original_sheet.iter_rows():
            for cell in row:
                assert result_sheet[cell.coordinate].value == cell.value

    source = original["Q1_单点组批"]
    target = result["Q1_单点组批"]
    assert [target.cell(1, column).value for column in range(1, 10)] == [
        source.cell(1, column).value for column in range(1, 10)
    ]
    populated = [
        row
        for row in target.iter_rows(min_row=2, max_col=9, values_only=True)
        if row[0] is not None
    ]
    assert len(populated) == 18
    box_ids = [
        box_id
        for row in populated
        for box_id in str(row[3]).split(";")
        if box_id
    ]
    assert len(box_ids) == 80
    assert len(set(box_ids)) == 80
    assert all(row[0].startswith("Q1-DP-") for row in populated)
    assert all(target.cell(1, column).style_id == source.cell(1, column).style_id for column in range(1, 10))
    assert [target.cell(2, column).number_format for column in range(5, 10)] == [
        "0.0",
        "0.000",
        "0.00",
        "0.000000",
        "0.0000",
    ]
